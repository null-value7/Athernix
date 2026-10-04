"""
Convierte ATHERNIX_EVALUACION.md a un documento Word (.docx) con formato
profesional (portada, titulos, tablas, listas, negritas), listo para
imprimir o entregar como documento de evaluacion institucional.

Uso:
    python scripts/md_to_docx_evaluacion.py
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Pt, Cm, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "ATHERNIX_EVALUACION.md"
DST = ROOT / "ATHERNIX_EVALUACION.docx"

COLOR_TITLE = RGBColor(0x1A, 0x1A, 0x2E)
COLOR_ACCENT = RGBColor(0x4F, 0x46, 0xE5)  # morado/indigo, coherente con la marca Athernix
COLOR_MUTED = RGBColor(0x55, 0x55, 0x55)


INLINE_RE = re.compile(r"(\*\*.+?\*\*|`.+?`|\*[^*\n]+?\*)")


def add_bold_runs(paragraph, text):
    """Agrega texto a un parrafo interpretando **negritas**, *cursivas* y `codigo` inline."""
    parts = INLINE_RE.split(text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.color.rgb = RGBColor(0xC7, 0x25, 0x4E)
        elif part.startswith("*") and part.endswith("*"):
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        else:
            run = paragraph.add_run(part)


def set_cell_shading(cell, color_hex):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color_hex)
    cell._tc.get_or_add_tcPr().append(shd)


def set_col_widths(table, widths_cm):
    for row in table.rows:
        for idx, width in enumerate(widths_cm):
            row.cells[idx].width = Cm(width)


def build_table(doc, header, rows):
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(header):
        hdr_cells[i].text = ""
        p = hdr_cells[i].paragraphs[0]
        run = p.add_run(h)
        run.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        set_cell_shading(hdr_cells[i], "4F46E5")
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            add_bold_runs(p, val)
    return table


def parse_table_block(lines, start):
    """lines[start] es la fila de encabezado de una tabla markdown."""
    header = [c.strip() for c in lines[start].strip().strip("|").split("|")]
    i = start + 2  # salta encabezado y la fila separadora ---|---
    rows = []
    while i < len(lines) and lines[i].strip().startswith("|"):
        row = [c.strip() for c in lines[i].strip().strip("|").split("|")]
        rows.append(row)
        i += 1
    return header, rows, i


def style_heading(paragraph, level):
    for run in paragraph.runs:
        run.font.color.rgb = COLOR_TITLE if level <= 2 else COLOR_ACCENT


def main():
    text = SRC.read_text(encoding="utf-8")
    lines = text.splitlines()

    doc = Document()

    # ── Estilos base del documento ──────────────────────────────────
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    for section in doc.sections:
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.2)
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)

    # ── Portada ──────────────────────────────────────────────────────
    titulo = doc.add_paragraph()
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = titulo.add_run("ATHERNIX")
    run.bold = True
    run.font.size = Pt(34)
    run.font.color.rgb = COLOR_ACCENT

    subtitulo = doc.add_paragraph()
    subtitulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitulo.add_run("Evaluación de Proyecto y Checklist de Presentación")
    run.font.size = Pt(16)
    run.font.color.rgb = COLOR_TITLE

    doc.add_paragraph()
    nota = doc.add_paragraph()
    nota.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = nota.add_run(
        "Documento generado a partir de una auditoría real del repositorio "
        "(estructura, dependencias, scripts, README y guion de presentación), "
        "comparado con plataformas EdTech VR de referencia "
        "(Labster, ENGAGE VR, VictoryXR, Google Arts & Culture)."
    )
    run.italic = True
    run.font.size = Pt(10.5)
    run.font.color.rgb = COLOR_MUTED

    doc.add_page_break()

    # ── Cuerpo: recorrido linea por linea del markdown ───────────────
    def is_block_start(s):
        return bool(re.match(r"^(#{1,4}\s|>|\-\s|\d+\.\s|\|)", s)) or s == ""

    def gather_continuation(start_idx):
        """Une las lineas siguientes que son continuacion del mismo parrafo/item
        (markdown hace 'soft wrap': una idea sigue en la linea de abajo sin blanco)."""
        buf = []
        j = start_idx
        while j < n and lines[j].strip() != "" and not is_block_start(lines[j].strip()):
            buf.append(lines[j].strip())
            j += 1
        return buf, j

    i = 0
    n = len(lines)
    while i < n:
        line = lines[i]
        stripped = line.strip()

        if stripped == "" or stripped == "---":
            i += 1
            continue

        # Cita / nota introductoria ("> ...")
        if stripped.startswith(">"):
            p = doc.add_paragraph(style="Intense Quote")
            textos = [stripped.lstrip(">").strip()]
            i += 1
            while i < n and lines[i].strip().startswith(">"):
                textos.append(lines[i].strip().lstrip(">").strip())
                i += 1
            add_bold_runs(p, " ".join(textos))
            continue

        # Encabezados
        m = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            texto = m.group(2).replace("✅ ", "")
            heading = doc.add_heading(level=min(level, 3))
            add_bold_runs(heading, texto)
            style_heading(heading, level)
            i += 1
            continue

        # Tabla markdown
        if stripped.startswith("|") and i + 1 < n and re.match(r"^\|?[\s:-]+\|", lines[i + 1].strip()):
            header, rows, i = parse_table_block(lines, i)
            build_table(doc, header, rows)
            doc.add_paragraph()
            continue

        # Lista numerada (incluye las indentadas, p.ej. sub-listas de objetivos especificos)
        m = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if m:
            extra, i = gather_continuation(i + 1)
            p = doc.add_paragraph(style="List Number")
            add_bold_runs(p, " ".join([m.group(2)] + extra))
            continue

        # Lista con guiones
        if stripped.startswith("- "):
            extra, i = gather_continuation(i + 1)
            p = doc.add_paragraph(style="List Bullet")
            add_bold_runs(p, " ".join([stripped[2:]] + extra))
            continue

        # Parrafo normal (puede continuar en la siguiente linea sin romper)
        extra, i = gather_continuation(i + 1)
        p = doc.add_paragraph()
        add_bold_runs(p, " ".join([stripped] + extra))

    doc.save(DST)
    print(f"Documento generado: {DST}")


if __name__ == "__main__":
    main()
